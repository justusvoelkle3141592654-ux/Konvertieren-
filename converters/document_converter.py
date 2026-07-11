"""Dokument-Konvertierungen: PDF, DOCX, TXT.

- PDF → DOCX:  pdf2docx (layout-treu)
- PDF → TXT:   pypdf (Textextraktion)
- DOCX → TXT:  python-docx
- DOCX → PDF:  LibreOffice (Linux/überall) oder docx2pdf (Word auf Windows/Mac)
- TXT → PDF:   fpdf2
- TXT → DOCX:  python-docx
"""

import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from converters.errors import ConversionError


def convert(data: bytes, src_ext: str, dst_ext: str) -> bytes:
    handlers = {
        ("pdf", "docx"): _pdf_to_docx,
        ("pdf", "txt"): _pdf_to_txt,
        ("docx", "txt"): _docx_to_txt,
        ("docx", "pdf"): _docx_to_pdf,
        ("txt", "pdf"): _txt_to_pdf,
        ("txt", "docx"): _txt_to_docx,
    }
    handler = handlers.get((src_ext, dst_ext))
    if handler is None:
        raise ConversionError(
            f"Die Umwandlung von .{src_ext} nach .{dst_ext} wird nicht unterstützt."
        )
    return handler(data)


def _decode_text(data: bytes) -> str:
    """Dekodiert Textdateien robust (UTF-8 zuerst, dann Latin-1 als Fallback)."""
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ConversionError("Die Textdatei konnte nicht dekodiert werden.")


def _pdf_to_docx(data: bytes) -> bytes:
    try:
        from pdf2docx import Converter
    except ImportError:
        raise ConversionError(
            "Das Paket 'pdf2docx' ist nicht installiert "
            "(pip install pdf2docx)."
        )

    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = Path(tmp) / "input.pdf"
        docx_path = Path(tmp) / "output.docx"
        pdf_path.write_bytes(data)
        try:
            converter = Converter(str(pdf_path))
            try:
                converter.convert(str(docx_path))
            finally:
                converter.close()
        except Exception as exc:
            raise ConversionError(
                f"PDF konnte nicht in DOCX umgewandelt werden: {exc}"
            )
        if not docx_path.exists():
            raise ConversionError("PDF → DOCX hat keine Ausgabedatei erzeugt.")
        return docx_path.read_bytes()


def _pdf_to_txt(data: bytes) -> bytes:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ConversionError(
            "Das Paket 'pypdf' ist nicht installiert (pip install pypdf)."
        )

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ConversionError(f"PDF konnte nicht gelesen werden: {exc}")

    text = "\n\n".join(pages).strip()
    if not text:
        raise ConversionError(
            "In diesem PDF wurde kein Text gefunden. "
            "Gescannte PDFs (nur Bilder) benötigen eine OCR-Software."
        )
    return text.encode("utf-8")


def _docx_to_txt(data: bytes) -> bytes:
    try:
        from docx import Document
    except ImportError:
        raise ConversionError(
            "Das Paket 'python-docx' ist nicht installiert "
            "(pip install python-docx)."
        )

    try:
        doc = Document(io.BytesIO(data))
        parts = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
    except Exception as exc:
        raise ConversionError(f"DOCX konnte nicht gelesen werden: {exc}")

    return "\n".join(parts).encode("utf-8")


def _docx_to_pdf(data: bytes) -> bytes:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")

    with tempfile.TemporaryDirectory() as tmp:
        docx_path = Path(tmp) / "input.docx"
        pdf_path = Path(tmp) / "input.pdf"
        docx_path.write_bytes(data)

        if soffice:
            try:
                subprocess.run(
                    [
                        soffice, "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, str(docx_path),
                    ],
                    capture_output=True,
                    timeout=120,
                    check=True,
                )
            except subprocess.TimeoutExpired:
                raise ConversionError(
                    "LibreOffice hat zu lange gebraucht (Timeout nach 2 Minuten)."
                )
            except subprocess.CalledProcessError as exc:
                stderr = (exc.stderr or b"").decode(errors="replace").strip()
                raise ConversionError(
                    f"LibreOffice konnte das DOCX nicht umwandeln: {stderr}"
                )
            if pdf_path.exists():
                return pdf_path.read_bytes()
            # soffice meldet auch bei Fehlern oft Exit-Code 0 – z.B. wenn nur
            # libreoffice-core ohne die Writer-Komponente installiert ist
            raise ConversionError(
                "LibreOffice hat keine PDF-Datei erzeugt. Häufige Ursache: "
                "die Writer-Komponente fehlt "
                "(Linux: sudo apt install libreoffice-writer)."
            )

        # Fallback: docx2pdf (benötigt Microsoft Word, nur Windows/macOS)
        if sys.platform in ("win32", "darwin"):
            try:
                from docx2pdf import convert as docx2pdf_convert
            except ImportError:
                raise ConversionError(
                    "Für DOCX → PDF wird LibreOffice oder das Paket 'docx2pdf' "
                    "(mit Microsoft Word) benötigt."
                )
            try:
                docx2pdf_convert(str(docx_path), str(pdf_path))
            except Exception as exc:
                raise ConversionError(
                    f"docx2pdf (Microsoft Word) fehlgeschlagen: {exc}"
                )
            if pdf_path.exists():
                return pdf_path.read_bytes()
            raise ConversionError("docx2pdf hat keine PDF-Datei erzeugt.")

        raise ConversionError(
            "Für DOCX → PDF wird LibreOffice benötigt. "
            "Installation z.B. mit: sudo apt install libreoffice"
        )


def _txt_to_pdf(data: bytes) -> bytes:
    try:
        from fpdf import FPDF
    except ImportError:
        raise ConversionError(
            "Das Paket 'fpdf2' ist nicht installiert (pip install fpdf2)."
        )

    text = _decode_text(data)
    try:
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        # Helvetica kann nur Latin-1: nicht darstellbare Zeichen ersetzen.
        # Ein einziger multi_cell-Aufruf verarbeitet alle Zeilenumbrüche und
        # vermeidet Cursor-Probleme neuerer fpdf2-Versionen.
        safe_text = text.encode("latin-1", errors="replace").decode("latin-1")
        pdf.multi_cell(0, 6, safe_text or " ")
        return bytes(pdf.output())
    except Exception as exc:
        raise ConversionError(f"TXT konnte nicht in PDF umgewandelt werden: {exc}")


def _txt_to_docx(data: bytes) -> bytes:
    try:
        from docx import Document
    except ImportError:
        raise ConversionError(
            "Das Paket 'python-docx' ist nicht installiert "
            "(pip install python-docx)."
        )

    text = _decode_text(data)
    try:
        doc = Document()
        for line in text.splitlines() or [""]:
            doc.add_paragraph(line)
        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()
    except Exception as exc:
        raise ConversionError(f"TXT konnte nicht in DOCX umgewandelt werden: {exc}")
